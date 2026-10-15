"""Privacy UI with the public build, real identity, API and PostgreSQL."""

import importlib.util
import json
import os
import secrets
import sys
from io import BytesIO
from pathlib import Path
from uuid import uuid4

import pytest
from openpyxl import load_workbook
from playwright.async_api import async_playwright, expect
from sqlalchemy import text

_spec = importlib.util.spec_from_file_location(
    "frontend_privacy_core", Path(__file__).with_name("test_core_contract.py")
)
core = importlib.util.module_from_spec(_spec)
sys.modules[_spec.name] = core
_spec.loader.exec_module(core)
banco = core.banco
harness = core.harness
backend = core.backend


@pytest.mark.parametrize("browser_name", ["chromium", "firefox", "webkit"])
@pytest.mark.parametrize(
    "viewport", [{"width": 390, "height": 844}, {"width": 1440, "height": 900}]
)
async def test_public_spa_privacidade_export_isolation_close(
    harness, engine_admin, viewport, browser_name
):
    await harness.restart()
    trace = (
        Path(os.environ["CORE_TRACE_DIR"])
        / f"privacidade-{browser_name}-{viewport['width']}.zip"
    )
    async with async_playwright() as playwright:
        browser = await getattr(playwright, browser_name).launch()
        contexts = [
            await browser.new_context(viewport=viewport, accept_downloads=True)
            for _ in range(2)
        ]
        pages = [await context.new_page() for context in contexts]
        users = []
        for page in pages:
            session, _ = await backend.register(
                page,
                uuid4().hex + "@example.org",
                secrets.token_urlsafe(24),
                harness=harness,
            )
            users.append(session["usuario_id"])
        uid, other = users
        await core.grant(pages[0], engine_admin, uid)
        names = ["Registros próprios " + uuid4().hex, "Outra pessoa " + uuid4().hex]
        async with engine_admin.begin() as conn:
            for owner, name in zip(users, names, strict=True):
                await conn.execute(
                    text("INSERT INTO bancas(usuario_id,nome) VALUES (:uid,:name)"),
                    {"uid": owner, "name": name},
                )
            await conn.execute(
                text(
                    "UPDATE assinaturas SET trial_ends_at=clock_timestamp()-interval '1 second' WHERE usuario_id=:uid"
                ),
                {"uid": uid},
            )
        assert (await core.request(pages[0], "/api/v1/billing/status"))["data"][
            "access"
        ] == "READ_ONLY"
        # Cookie writes require CSRF, including the explicit commercial exception.
        assert (await core.request(pages[0], "/api/v1/usuario/me", method="DELETE"))[
            "status"
        ] == 403
        page = pages[0]
        writes = []
        page.on(
            "request",
            lambda request: (
                writes.append(request)
                if request.method == "DELETE"
                and request.url.endswith("/api/v1/usuario/me")
                else None
            ),
        )
        await contexts[0].tracing.start(
            screenshots=False, snapshots=False, sources=False
        )
        try:
            await page.goto(
                backend.FRONT + "/configuracoes/privacidade?casa=7&apagadas=1"
            )
            await expect(
                page.get_by_role("heading", name="Privacidade", exact=True)
            ).to_be_visible()
            await expect(
                page.get_by_role("link", name="Voltar às configurações")
            ).to_have_attribute("href", "/configuracoes?casa=7&apagadas=1")
            for label, extension in (("JSON", "json"), ("Excel", "xlsx")):
                async with page.expect_download() as download_info:
                    await page.get_by_role(
                        "button", name="Baixar " + label, exact=True
                    ).click()
                download = await download_info.value
                assert download.suggested_filename == "meus-dados." + extension
                content = Path(await download.path()).read_bytes()
                if extension == "json":
                    exported = json.loads(content)
                    assert exported["usuario"]["id"] == uid
                    assert any(row["nome"] == names[0] for row in exported["bancas"])
                    assert names[1] not in content.decode()
                    assert (
                        "token_hash" not in content.decode()
                        and "provider_customer_ref" not in content.decode()
                    )
                else:
                    book = load_workbook(BytesIO(content), read_only=True)
                    try:
                        assert {"usuario", "bancas", "eventos"} <= set(book.sheetnames)
                        rows = list(book["bancas"].values)
                        assert any(names[0] in row for row in rows)
                        assert not any(names[1] in row for row in rows)
                    finally:
                        book.close()
            await page.get_by_role(
                "button", name="Solicitar encerramento", exact=True
            ).click()
            modal = page.get_by_role("dialog")
            await expect(
                modal.get_by_role("button", name="Manter minha conta")
            ).to_be_focused()
            await expect(
                modal.get_by_role("button", name="Confirmar encerramento")
            ).to_be_disabled()
            await page.keyboard.press("Escape")
            assert len(writes) == 0
            # Synthetic imported-source marker exercises the real 409 guard.
            async with engine_admin.begin() as conn:
                raw_id = await conn.scalar(
                    text(
                        "INSERT INTO uploads(usuario_id,filename,chat_id) VALUES (:uid,'privacy-test.zip',9876) RETURNING id"
                    ),
                    {"uid": uid},
                )
            await page.get_by_role(
                "button", name="Solicitar encerramento", exact=True
            ).click()
            await modal.get_by_role("checkbox").check()
            await modal.get_by_role("button", name="Confirmar encerramento").click()
            await expect(
                page.get_by_role(
                    "heading", name="O encerramento precisa de acompanhamento"
                )
            ).to_be_visible()
            assert len(writes) == 1
            async with engine_admin.begin() as conn:
                assert await conn.scalar(
                    text("SELECT ativo FROM usuarios WHERE id=:uid"), {"uid": uid}
                )
                await conn.execute(
                    text("DELETE FROM uploads WHERE id=:id AND usuario_id=:uid"),
                    {"id": raw_id, "uid": uid},
                )
            await page.get_by_role(
                "button", name="Solicitar encerramento", exact=True
            ).click()
            await modal.get_by_role("checkbox").check()
            await modal.get_by_role("button", name="Confirmar encerramento").click()
            await expect(
                page.get_by_role("heading", name="Entrar", exact=True)
            ).to_be_visible()
            await expect(
                page.get_by_text("Conta desativada e anonimizada.", exact=False)
            ).to_be_visible()
            assert len(writes) == 2
            assert (await core.request(page, "/api/v1/usuario/me/export"))[
                "status"
            ] == 401
            async with engine_admin.connect() as conn:
                assert not await conn.scalar(
                    text("SELECT ativo FROM usuarios WHERE id=:uid"), {"uid": uid}
                )
                assert (
                    await conn.scalar(
                        text("SELECT count(*) FROM bancas WHERE usuario_id=:uid"),
                        {"uid": uid},
                    )
                    == 0
                )
                assert await conn.scalar(
                    text("SELECT ativo FROM usuarios WHERE id=:uid"), {"uid": other}
                )
            other_export = await core.request(pages[1], "/api/v1/usuario/me/export")
            assert other_export["status"] == 200
            assert other_export["data"]["usuario"]["id"] == other
            assert any(
                row["nome"] == names[1] for row in other_export["data"]["bancas"]
            )
        finally:
            await contexts[0].tracing.stop(path=trace)
            core.sanitize_trace(trace)
            for context in contexts:
                await context.close()
            await browser.close()
