"""Real cookie contract from the public SPA; no session override enters its bundle.

Run with the pinned backend as cwd, after building the frontend. Its sandbox fixtures
own disposable PostgreSQL/OIDC/SMTP. This server stands in for the same-origin proxy.
"""

import asyncio
import json
import mimetypes
import os
import re
import secrets
import time
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qsl, urlencode, urlsplit
from uuid import uuid4

import httpx
import pytest
from playwright.async_api import Error as BrowserError
from playwright.async_api import async_playwright, expect
from sqlalchemy import text
from sqlalchemy.engine import make_url
from sqlalchemy.ext.asyncio import create_async_engine
from tests.conftest import _preparar

from tests.identity import journey as backend

ROOT = Path(__file__).resolve().parents[2]
DIST = ROOT / "dist"


def serve_public_build(directory=DIST, security_directory=ROOT / "dist-security"):
    security = (security_directory / "default.conf").read_text()
    csp = re.search(r'add_header Content-Security-Policy "([^"]+)"', security).group(1)
    assert "connect-src 'self'" in csp and "unsafe-inline" not in csp

    class PublicBuild(BaseHTTPRequestHandler):
        def log_message(self, *_args):
            pass  # Cookie, CSRF, protocol URLs and personal data stay out of logs.

        def handle_request(self):
            path = urlsplit(self.path).path
            if path.startswith(("/auth/", "/api/")):
                body = self.rfile.read(int(self.headers.get("content-length", "0")))
                headers = {
                    k: v
                    for k, v in self.headers.items()
                    if k.lower() not in {"host", "connection", "accept-encoding"}
                }
                with httpx.Client(timeout=15, follow_redirects=False) as client:
                    result = client.request(
                        self.command,
                        backend.API + self.path,
                        headers=headers,
                        content=body,
                    )
                self.send_response(result.status_code)
                for key, value in result.headers.multi_items():
                    if key.lower() not in {
                        "transfer-encoding",
                        "content-length",
                        "connection",
                        "content-encoding",
                    }:
                        self.send_header(key, value)
                self.send_header("Content-Length", str(len(result.content)))
                self.end_headers()
                self.wfile.write(result.content)
                return
            if self.command != "GET":
                self.send_error(405)
                return
            if path == "/config.json":
                content = json.dumps({"VITE_API_URL": backend.FRONT}).encode()
                content_type = "application/json"
            else:
                target = (directory / path.lstrip("/")).resolve()
                if not target.is_relative_to(directory.resolve()):
                    self.send_error(404)
                    return
                if path.startswith(("/assets/", "/fontes/")) and not target.is_file():
                    self.send_error(404)
                    return
                if not target.is_file():
                    target = directory / "index.html"
                content = target.read_bytes()
                content_type = (
                    mimetypes.guess_type(target)[0] or "application/octet-stream"
                )
                if target.suffix == ".js":
                    content_type = "application/javascript"
            self.send_response(200)
            self.send_header("Content-Type", content_type)
            self.send_header("Content-Security-Policy", csp)
            self.send_header("Cache-Control", "no-store")
            self.send_header("Content-Length", str(len(content)))
            self.end_headers()
            self.wfile.write(content)

        do_GET = handle_request
        do_POST = handle_request
        do_DELETE = handle_request

    return ThreadingHTTPServer(("127.0.0.1", 58001), PublicBuild)


async def browser_request(page, path, *, method="GET", csrf=None, body=None):
    return await page.evaluate(
        """async ({path,method,csrf,body}) => {
          const headers={}; if(csrf) headers['X-CSRF-Token']=csrf;
          if(body) headers['Content-Type']='application/json';
          const response=await fetch(path, {method, credentials:'include', headers,
            cache:'no-store', body:body ? JSON.stringify(body) : undefined});
          return {status:response.status, data:await response.json(),
            headers:Object.fromEntries(response.headers)};
        }""",
        {"path": path, "method": method, "csrf": csrf, "body": body},
    )


async def complete_hosted_entry(page, email, password, destination):
    # Keycloak can remember a different SSO account and show password-only. The
    # hosted "reset-login" action selects another account without altering cookies.
    # Source: keycloak/keycloak 26.7.4, base/login/template.ftl.
    try:
        await page.locator("#kc-form-login").wait_for()
        if not await page.locator("#username").is_visible():
            await page.locator("#reset-login").click()
        await page.locator("#username").fill(email)
        if not await page.locator("#password").count():
            await page.locator("#kc-login").click()
        await page.locator("#password").fill(password)
        await page.locator("#kc-login").click()
        await page.wait_for_url(destination, timeout=30000)
    except BrowserError:
        # Playwright diagnostics can contain protocol URLs or filled credentials.
        pytest.fail("Hosted account entry did not complete", pytrace=False)


@pytest.fixture
async def banco():
    # Each harness generates its own keyring. Never reuse encrypted rows from another
    # sandbox harness: its revocation outbox must remain with the keys that created it.
    url = make_url(os.environ["TEST_DATABASE_URL"])
    assert url.host == "127.0.0.1" and url.port == 55432
    assert url.database == "bancaemdia_identity", (
        "Only the disposable CI database is allowed"
    )
    name = "identity_frontend_" + uuid4().hex
    admin = create_async_engine(
        url.set(database="postgres"), isolation_level="AUTOCOMMIT"
    )
    try:
        async with admin.connect() as conn:
            await conn.execute(text(f"CREATE DATABASE {name}"))
        isolated_url = url.set(database=name).render_as_string(hide_password=False)
        yield await asyncio.to_thread(_preparar, isolated_url)
    finally:
        async with admin.connect() as conn:
            await conn.execute(text(f"DROP DATABASE {name}"))
        await admin.dispose()


@pytest.fixture
async def harness(banco, engine_admin, tmp_path, monkeypatch):
    assert (DIST / "index.html").is_file(), "Build the public SPA before acceptance"
    monkeypatch.setattr(backend, "serve_frontend", serve_public_build)
    monkeypatch.setattr(backend, "fetch", browser_request)
    async for instance in backend.harness.__wrapped__(banco, engine_admin, tmp_path):
        yield instance


@pytest.fixture
async def harness_filtros(banco, engine_admin, tmp_path, monkeypatch):
    directory = ROOT / "dist-filtros-fixture"
    assert (directory / "index.html").is_file(), "Build the isolated exercise"
    monkeypatch.setattr(
        backend,
        "serve_frontend",
        lambda: serve_public_build(directory, ROOT / "dist-filtros-security"),
    )
    monkeypatch.setattr(backend, "fetch", browser_request)
    async for instance in backend.harness.__wrapped__(banco, engine_admin, tmp_path):
        yield instance


@pytest.mark.parametrize(
    "viewport", [{"width": 390, "height": 844}, {"width": 1440, "height": 900}]
)
async def test_authenticated_filter_components_real_cookie(
    harness_filtros, engine_admin, viewport
):
    """Actual components and disposable input; session, SQL and results are real."""
    import io

    from bancaemdia import models
    from openpyxl import load_workbook
    from sqlalchemy.ext.asyncio import AsyncSession

    async with async_playwright() as playwright:
        browser = await playwright.chromium.launch()
        contexts = [await browser.new_context(viewport=viewport) for _ in range(2)]
        pages = [await context.new_page() for context in contexts]
        sessions = []
        violations = []
        for page in pages:

            def page_error(_error, current=page):
                if urlsplit(current.url).netloc == urlsplit(backend.FRONT).netloc:
                    violations.append("filter exercise pageerror")

            page.on("pageerror", page_error)
            await page.add_init_script(
                "if (location.origin === "
                + json.dumps(backend.FRONT)
                + ") { document.addEventListener('securitypolicyviolation', () => { window.__filterCspFailure=true; }); }"
            )
            session, _ = await backend.register(
                page,
                uuid4().hex + "@example.org",
                secrets.token_urlsafe(24),
                harness=harness_filtros,
            )
            sessions.append(session)
        created = await browser_request(
            pages[0],
            "/api/v1/apostas",
            method="POST",
            csrf=sessions[0]["csrf_token"],
            body={
                "casa": "betano",
                "odd": 2,
                "stake_unidades": 1,
                "data_aposta": "2026-10-06T12:00:00-03:00",
            },
        )
        assert created["status"] == 201
        key = created["data"]["aposta"]["chave"]
        historic_id = 9007199254740993
        async with AsyncSession(engine_admin, expire_on_commit=False) as db:
            group = models.GrupoAposta(
                id=historic_id,
                usuario_id=sessions[0]["usuario_id"],
                nome="Grupo histórico exato",
                arquivado=True,
            )
            db.add(group)
            await db.flush()
            aposta = await db.scalar(
                text("SELECT id FROM apostas WHERE chave=:key AND usuario_id=:user"),
                {"key": key, "user": sessions[0]["usuario_id"]},
            )
            db.add(
                models.ApostaGrupo(
                    usuario_id=sessions[0]["usuario_id"],
                    aposta_id=aposta,
                    grupo_id=historic_id,
                )
            )
            await db.commit()
        assert (
            await browser_request(
                pages[0],
                "/api/v1/apostas/" + key,
                method="DELETE",
                csrf=sessions[0]["csrf_token"],
            )
        )["status"] == 200
        query = "?grupo=9007199254740993&apagadas=1&estado=PENDENTE&origem=manual&desde=2026-10-06&ate=2026-10-06&page_size=1"
        reads = []

        def observe(request):
            if "/api/v1/filtros/" in request.url:
                reads.append(urlsplit(request.url).path)
                assert "authorization" not in request.headers

        pages[0].on("request", observe)
        await pages[0].goto(backend.FRONT + "/testes/filtros-autenticados" + query)
        await expect(
            pages[0].get_by_role("button", name="Grupo Grupo histórico exato (inativa)")
        ).to_be_visible()
        await expect(pages[0].get_by_label("Resposta da lista")).to_contain_text(key)

        async def output(page, label):
            await expect(page.get_by_label(label)).not_to_have_text("null")
            return json.loads(await page.get_by_label(label).inner_text())

        listed = await output(pages[0], "Resposta da lista")
        summary = await output(pages[0], "Resposta do resumo")
        assert listed["pagination"]["total"] == 1
        assert summary["resumo"]["total_apostas"] == 1
        assert summary["resumo"]["pendentes"] == 1
        assert summary["resumo"]["lucro_centavos"] == 0
        assert len(set(reads)) == 9
        params = await output(pages[0], "Consulta normalizada")
        assert (
            params["grupo_id"] == "9007199254740993"
            and params["visibilidade"] == "apagadas"
        )
        assert "incluir_apagadas" not in params and "periodo" not in params
        common = {k: v for k, v in params.items() if k not in {"page", "page_size"}}
        metrics = await browser_request(
            pages[0], "/api/v1/painel/filtrado/metricas?" + urlencode(common)
        )
        assert metrics["status"] == 200
        assert metrics["data"]["total_periodo"]["total_apostas"] == 1
        assert metrics["data"]["total_periodo"]["lucro_centavos"] == 0
        # XLSX is authoritative aggregates, not a client-side export of one page.
        exported = await contexts[0].request.get(
            backend.FRONT + "/api/v1/painel/filtrado/export?" + urlencode(common)
        )
        assert exported.status == 200
        workbook = load_workbook(io.BytesIO(await exported.body()))
        assert workbook.sheetnames == [
            "Resumo",
            "Por casa",
            "Por tipster",
            "Por mercado",
            "Por periodo",
            "Evolucao",
        ]
        rows = list(workbook["Resumo"].values)
        exported_summary = dict(zip(rows[0], rows[1]))
        assert exported_summary["total_apostas"] == 1
        assert exported_summary["pendentes"] == 1
        assert exported_summary["lucro_centavos"] == 0
        await pages[0].get_by_role("button", name="Próxima página", exact=True).click()
        await expect(pages[0].get_by_label("Resposta da lista")).to_contain_text(
            '"page":2'
        )
        listed = await output(pages[0], "Resposta da lista")
        assert listed["data"] == [] and listed["pagination"]["total"] == 1
        await pages[0].reload()
        await expect(
            pages[0].get_by_role("button", name="Grupo Grupo histórico exato (inativa)")
        ).to_be_visible()
        await pages[1].goto(backend.FRONT + "/testes/filtros-autenticados" + query)
        listed_other = await output(pages[1], "Resposta da lista")
        assert listed_other["data"] == [] and listed_other["pagination"]["total"] == 0
        await expect(
            pages[1].get_by_role(
                "button",
                name="Grupo Identificador 9007199254740993 (nome indisponível)",
            )
        ).to_be_visible()
        await pages[0].get_by_role("button", name="Remover filtro Visibilidade").click()
        await expect(pages[0].get_by_label("Resposta da lista")).to_contain_text(
            '"total":0'
        )
        await pages[0].get_by_role("link", name="Sair do exercício").click()
        await pages[0].get_by_role("button", name="Confirmar saída").click()
        await expect(
            pages[0].get_by_role("heading", name="Entrar", exact=True)
        ).to_be_visible()
        assert (await browser_request(pages[0], "/api/v1/filtros/grupos"))[
            "status"
        ] == 401
        for context, page in zip(contexts, pages):
            assert not await page.evaluate("window.__filterCspFailure === true")
            assert all(
                cookie["httpOnly"] for cookie in await context.cookies(backend.FRONT)
            )
            assert await page.evaluate("document.cookie") == ""
            assert (
                await page.evaluate(
                    "Object.keys(localStorage).length + Object.keys(sessionStorage).length"
                )
                == 0
            )
        assert not violations
        await browser.close()


@pytest.mark.parametrize(
    "viewport", [{"width": 390, "height": 844}, {"width": 1440, "height": 900}]
)
async def test_public_spa_real_identity_cookie_contract(harness, viewport):
    async with async_playwright() as playwright:
        browser = await playwright.chromium.launch()
        contexts = [await browser.new_context(viewport=viewport) for _ in range(2)]
        pages = [await context.new_page() for context in contexts]
        violations = []
        for page in pages:

            def public_page_error(_error, current_page=page):
                if urlsplit(current_page.url).netloc == urlsplit(backend.FRONT).netloc:
                    violations.append("public SPA pageerror")

            page.on("pageerror", public_page_error)
            # This monitor belongs to the SPA, not the hosted issuer's documents. It
            # also records early CSP violations before documentElement exists.
            await page.add_init_script(
                "if (location.origin === "
                + json.dumps(backend.FRONT)
                + ") { document.addEventListener('securitypolicyviolation',"
                + "() => { window.__identityCspFailure=true; }); }"
            )
        sessions = []
        for page in pages:
            email = uuid4().hex + "@example.org"
            session, _ = await backend.register(
                page, email, secrets.token_urlsafe(24), harness=harness
            )
            sessions.append(session)
            await page.goto(backend.FRONT + "/login")
            await expect(
                page.get_by_role("heading", name="Entrar", exact=True)
            ).to_be_visible()
            assert not await page.evaluate("window.__identityCspFailure === true")
            assert not {"access_token", "refresh_token", "id_token"}.intersection(
                session
            )
            assert await page.evaluate("Object.keys(localStorage).length") == 0
            assert await page.evaluate("Object.keys(sessionStorage).length") == 0
            if await page.evaluate("document.cookie"):
                pytest.fail("Session cookies must not be readable by scripts")
        assert sessions[0]["usuario_id"] != sessions[1]["usuario_id"]
        # Actual browser cookie authorization and CSRF; only disposable financial input.
        rejected = await browser_request(pages[0], "/auth/refresh", method="POST")
        assert rejected["status"] == 403 and rejected["data"]["code"] == "csrf_failed"
        renewed = await browser_request(
            pages[0], "/auth/refresh", method="POST", csrf=sessions[0]["csrf_token"]
        )
        assert renewed["status"] == 200
        live = renewed["data"]
        assert live["usuario_id"] == sessions[0]["usuario_id"]
        assert live["session_version"] != sessions[0]["session_version"]
        if live["csrf_token"] == sessions[0]["csrf_token"]:
            pytest.fail("Refresh must rotate the CSRF proof")
        created = await browser_request(
            pages[0],
            "/api/v1/apostas",
            method="POST",
            csrf=live["csrf_token"],
            body={
                "casa": "betano",
                "odd": 2,
                "stake_unidades": 1,
                "data_aposta": "2026-10-06T12:00:00-03:00",
            },
        )
        assert created["status"] == 201
        key = created["data"]["aposta"]["chave"]
        assert (await browser_request(pages[0], "/api/v1/apostas/" + key))[
            "status"
        ] == 200
        assert (await browser_request(pages[1], "/api/v1/apostas/" + key))[
            "status"
        ] == 404
        # #17 uses pages and inclusive/exclusive instants; never filters a page locally.
        selected = await browser_request(
            pages[0],
            "/api/v1/apostas?estado=PENDENTE&page=1&page_size=1"
            "&desde=2026-10-06T03%3A00%3A00Z&ate=2026-10-07T03%3A00%3A00Z",
        )
        assert selected["status"] == 200
        assert selected["data"]["pagination"] == {"page": 1, "page_size": 1, "total": 1}
        assert selected["data"]["data"][0]["chave"] == key
        for query in (
            "estado=RED",
            "desde=2026-10-07T03%3A00%3A00Z",
            "ate=2026-10-06T15%3A00%3A00Z",
        ):
            empty = await browser_request(pages[0], "/api/v1/apostas?" + query)
            assert empty["status"] == 200 and empty["data"]["pagination"]["total"] == 0
        exact_start = await browser_request(
            pages[0], "/api/v1/apostas?desde=2026-10-06T15%3A00%3A00Z"
        )
        assert (
            exact_start["status"] == 200
            and exact_start["data"]["pagination"]["total"] == 1
        )
        second = await browser_request(pages[0], "/api/v1/apostas?page=2&page_size=1")
        assert (
            second["data"]["data"] == [] and second["data"]["pagination"]["total"] == 1
        )
        other = await browser_request(pages[1], "/api/v1/apostas?estado=PENDENTE")
        assert other["status"] == 200 and other["data"]["pagination"]["total"] == 0
        removed = await browser_request(
            pages[0], "/api/v1/apostas/" + key, method="DELETE", csrf=live["csrf_token"]
        )
        assert removed["status"] == 200
        active = await browser_request(pages[0], "/api/v1/apostas")
        included = await browser_request(
            pages[0], "/api/v1/apostas?incluir_apagadas=true"
        )
        assert active["data"]["pagination"]["total"] == 0
        assert included["data"]["pagination"]["total"] == 1
        restored = await browser_request(
            pages[0],
            "/api/v1/apostas/" + key + "/restaurar",
            method="POST",
            csrf=live["csrf_token"],
        )
        assert restored["status"] == 200
        # The actual provider now owns cookie/CSRF lookup, clearing and confirmation.
        await pages[0].goto(backend.FRONT + "/sair")
        await pages[0].get_by_role("button", name="Confirmar saída").click()
        await expect(
            pages[0].get_by_role("heading", name="Entrar", exact=True)
        ).to_be_visible()
        assert (await browser_request(pages[0], "/auth/session"))["status"] == 401
        assert (await browser_request(pages[0], "/api/v1/apostas"))["status"] == 401
        assert (await browser_request(pages[1], "/auth/session"))["status"] == 200
        cache_control = (await browser_request(pages[1], "/auth/session"))["headers"][
            "cache-control"
        ]
        assert "no-store" in cache_control
        for context, page in zip(contexts, pages):
            assert not await page.evaluate("window.__identityCspFailure === true")
            assert all(c["httpOnly"] for c in await context.cookies(backend.FRONT))
            await page.goto(backend.FRONT + "/painel")
            await expect(
                page.get_by_role(
                    "heading",
                    name="Entrar" if page == pages[0] else "Painel",
                    exact=True,
                )
            ).to_be_visible()
            assert not await page.evaluate("window.__identityCspFailure === true")
        assert not violations
        await browser.close()


@pytest.mark.parametrize(
    "viewport", [{"width": 390, "height": 844}, {"width": 1440, "height": 900}]
)
async def test_public_spa_session_lifecycle(harness, engine_admin, viewport):
    """Real provider, Web Locks across tabs, cache switch and hosted return; no mock."""
    import jwt

    async with async_playwright() as playwright:
        browser = await playwright.chromium.launch()
        a, b = [await browser.new_context(viewport=viewport) for _ in range(2)]
        pa, pb = await a.new_page(), await b.new_page()
        email_a, email_b = (uuid4().hex + "@example.org" for _ in range(2))
        password_a, password_b = (secrets.token_urlsafe(24) for _ in range(2))
        sa, _ = await backend.register(pa, email_a, password_a, harness=harness)
        sb, _ = await backend.register(pb, email_b, password_b, harness=harness)
        assert sa["usuario_id"] != sb["usuario_id"]
        live = (await browser_request(pa, "/auth/session"))["data"]
        created = await browser_request(
            pa,
            "/api/v1/apostas",
            method="POST",
            csrf=live["csrf_token"],
            body={"casa": "betano", "odd": 2, "stake_unidades": 1, "freebet": False},
        )
        assert created["status"] == 201
        # Add one disposable pending review, to distinguish the real per-user caches.
        async with engine_admin.begin() as conn:
            result = await conn.execute(
                text(
                    "INSERT INTO revisao_pendente (usuario_id, motivo) VALUES (:id, 'sandbox')"
                ),
                {"id": sa["usuario_id"]},
            )
            assert result.rowcount == 1
        await pa.goto(backend.FRONT + "/painel?apagadas=1#serie")
        await expect(
            pa.get_by_role("heading", name="Painel", exact=True)
        ).to_be_visible()
        await expect(
            pa.get_by_role("link", name=re.compile(r"Revisão.*1 pendências")).first
        ).to_be_visible()
        second = await a.new_page()
        await second.goto(backend.FRONT + "/painel")
        await expect(
            second.get_by_role("heading", name="Painel", exact=True)
        ).to_be_visible()
        _, credentials = await harness.credentials(a)
        expiration = jwt.decode(
            credentials["internal"], options={"verify_signature": False}
        )["exp"]
        await asyncio.sleep(max(0, expiration - time.time()) + 1)
        grants = harness.network.token_requests
        await asyncio.gather(
            pa.goto(backend.FRONT + "/painel?apagadas=1#serie"),
            second.goto(backend.FRONT + "/painel"),
        )
        for page in (pa, second):
            await expect(
                page.get_by_role("heading", name="Painel", exact=True)
            ).to_be_visible()
        assert harness.network.token_requests == grants + 1
        refresh_required = (await browser_request(pa, "/auth/session"))["data"][
            "refresh_required"
        ]
        assert refresh_required is False
        # Log in as B in one tab; the other mounted SPA must discard A's private cache.
        await second.goto(backend.FRONT + "/login?destino=%2Fpainel")
        await second.get_by_role("button", name="Entrar com minha conta").click()
        await complete_hosted_entry(
            second, email_b, password_b, backend.FRONT + "/painel"
        )
        await expect(
            pa.get_by_role("link", name=re.compile(r"Revisão.*1 pendências"))
        ).to_have_count(0)
        await expect(
            pa.get_by_role("heading", name="Painel", exact=True)
        ).to_be_visible()
        current_person = (await browser_request(pa, "/auth/session"))["data"][
            "usuario_id"
        ]
        assert current_person == sb["usuario_id"]
        assert (await browser_request(pa, "/api/v1/apostas"))["data"]["pagination"][
            "total"
        ] == 0
        await pa.goto(backend.FRONT + "/sair")
        await pa.get_by_role("button", name="Confirmar saída").click()
        await expect(
            pa.get_by_role("heading", name="Entrar", exact=True)
        ).to_be_visible()
        assert (await browser_request(second, "/auth/session"))["status"] == 401
        # A fresh hosted entry initiated by the real SPA preserves filters and hash.
        await pa.goto(backend.FRONT + "/login?destino=%2Fpainel%3Fapagadas%3D1%23serie")
        await pa.get_by_role("button", name="Entrar com minha conta").click()
        await complete_hosted_entry(
            pa, email_a, password_a, backend.FRONT + "/painel?apagadas=1#serie"
        )
        await expect(
            pa.get_by_role("heading", name="Painel", exact=True)
        ).to_be_visible()
        assert await pa.evaluate("Object.keys(sessionStorage).length") == 0
        assert await pa.evaluate("Object.keys(localStorage).length") == 0
        assert not await pa.evaluate("document.cookie")
        await browser.close()


async def start_public_registration(page, email, destination):
    from urllib.parse import urlencode

    await page.goto(
        backend.FRONT + "/criar-conta?" + urlencode({"destino": destination})
    )
    await page.get_by_role("button", name="Continuar para criar conta").click()
    await page.get_by_role("link", name="Register", exact=True).click()
    await page.locator("#email").fill(email)
    await page.locator("#firstName").fill("Disposable")
    await page.locator("#lastName").fill("Acceptance")
    assert await page.locator("#password").count() == 0
    await page.locator('input[type="submit"],button[type="submit"]').click()
    return await backend.mail_link(email)


@pytest.mark.parametrize(
    "viewport", [{"width": 390, "height": 844}, {"width": 1440, "height": 900}]
)
async def test_public_spa_account_registration_and_recovery(
    harness, engine_admin, viewport
):
    """Actual SPA actions, issuer forms, confirmation SMTP and recovery revocation."""
    from urllib.parse import urlencode

    try:
        async with async_playwright() as playwright:
            browser = await playwright.chromium.launch()
            original = await browser.new_context(viewport=viewport)
            recovering = await browser.new_context(viewport=viewport)
            repeated = await browser.new_context(viewport=viewport)
            page, recovery, repeat = (
                await original.new_page(),
                await recovering.new_page(),
                await repeated.new_page(),
            )
            email = uuid4().hex + "@example.org"
            password, changed = (secrets.token_urlsafe(24) for _ in range(2))
            destination = "/painel?apagadas=1#serie"
            confirmation, message = await start_public_registration(
                page, email, destination
            )
            async with engine_admin.connect() as conn:
                count = await conn.scalar(
                    text("SELECT count(*) FROM usuarios WHERE email=:email"),
                    {"email": email},
                )
                assert count == 0, "Unconfirmed signup must not provision an API user"
            await page.goto(confirmation)
            await page.locator("#password-new").fill(password)
            await page.locator("#password-confirm").fill(password)
            await page.locator('input[type="submit"],button[type="submit"]').click()
            await page.wait_for_url(backend.FRONT + destination)
            await expect(
                page.get_by_role("heading", name="Painel", exact=True)
            ).to_be_visible()
            first = await browser_request(page, "/auth/session")
            assert first["status"] == 200
            person = first["data"]["usuario_id"]
            async with engine_admin.connect() as conn:
                assert (
                    await conn.scalar(
                        text("SELECT count(*) FROM usuarios WHERE email=:email"),
                        {"email": email},
                    )
                    == 1
                )
            # An already consumed email action cannot create another API session.
            await repeat.goto(confirmation)
            await repeat.goto(backend.FRONT + "/confirmar-email")
            await expect(
                repeat.get_by_role("button", name="Continuar confirmação")
            ).to_be_enabled()
            assert (await browser_request(repeat, "/auth/session"))["status"] == 401
            # An interrupted recovery changes neither password nor existing local session.
            recover_url = (
                backend.FRONT + "/esqueci-senha?" + urlencode({"destino": destination})
            )
            await recovery.goto(recover_url)
            await recovery.get_by_role(
                "button", name="Continuar para recuperar senha"
            ).click()
            await recovery.get_by_role(
                "link", name="Forgot Password?", exact=True
            ).click()
            await recovery.goto(recover_url)
            assert (await browser_request(page, "/auth/session"))["status"] == 200
            await recovery.get_by_role(
                "button", name="Continuar para recuperar senha"
            ).click()
            await recovery.get_by_role(
                "link", name="Forgot Password?", exact=True
            ).click()
            await recovery.locator("#username").fill(email)
            await recovery.locator('input[type="submit"],button[type="submit"]').click()
            reset, _ = await backend.mail_link(email, {message})
            await recovery.goto(reset)
            await recovery.locator("#password-new").fill(changed)
            await recovery.locator("#password-confirm").fill(changed)
            await recovery.locator('input[type="submit"],button[type="submit"]').click()
            await recovery.wait_for_url(backend.FRONT + destination)
            await expect(
                recovery.get_by_role("heading", name="Painel", exact=True)
            ).to_be_visible()
            recovered = await browser_request(recovery, "/auth/session")
            assert recovered["status"] == 200
            assert recovered["data"]["usuario_id"] == person
            assert (await browser_request(page, "/auth/session"))["status"] == 401
            await page.goto(backend.FRONT + "/painel")
            await expect(
                page.get_by_role("heading", name="Entrar", exact=True)
            ).to_be_visible()
            for current in (page, recovery, repeat):
                assert await current.evaluate("Object.keys(sessionStorage).length") == 0
                assert await current.evaluate("Object.keys(localStorage).length") == 0
                assert not await current.evaluate("document.cookie")
                assert not {"code", "state", "key", "token"}.intersection(
                    dict(parse_qsl(urlsplit(current.url).query))
                )
            await browser.close()
    except BrowserError:
        pytest.fail("Public account journey did not complete", pytrace=False)


@pytest.fixture
async def short_email_lifetime():
    """Only this disposable issuer; restore its setting even on a rejected action."""
    # Keycloak 26.7.4 VerifyEmail.java reads the realm's user-action lifetime.
    async with httpx.AsyncClient(timeout=15) as client:
        token = await client.post(
            "http://127.0.0.1:58080/realms/master/protocol/openid-connect/token",
            data={
                "grant_type": "password",
                "client_id": "admin-cli",
                "username": os.environ["KC_BOOTSTRAP_ADMIN_USERNAME"],
                "password": os.environ["KC_BOOTSTRAP_ADMIN_PASSWORD"],
            },
        )
        assert token.status_code == 200, "Sandbox issuer administration unavailable"
        headers = {"Authorization": "Bearer " + token.json()["access_token"]}
        endpoint = "http://127.0.0.1:58080/admin/realms/bancaemdia-acceptance"
        realm = await client.get(endpoint, headers=headers)
        assert realm.status_code == 200
        previous = realm.json()["actionTokenGeneratedByUserLifespan"]
        result = await client.put(
            endpoint, headers=headers, json={"actionTokenGeneratedByUserLifespan": 1}
        )
        assert result.status_code == 204
        try:
            yield
        finally:
            restored = await client.put(
                endpoint,
                headers=headers,
                json={"actionTokenGeneratedByUserLifespan": previous},
            )
            assert restored.status_code == 204


@pytest.mark.parametrize(
    "viewport", [{"width": 390, "height": 844}, {"width": 1440, "height": 900}]
)
async def test_public_spa_expired_confirmation(
    harness, engine_admin, short_email_lifetime, viewport
):
    """Real expired issuer mail, not a modified/synthetic token or a frozen clock."""
    try:
        async with async_playwright() as playwright:
            browser = await playwright.chromium.launch()
            context = await browser.new_context(viewport=viewport)
            page = await context.new_page()
            email = uuid4().hex + "@example.org"
            confirmation, _ = await start_public_registration(page, email, "/painel")
            await asyncio.sleep(2)
            await page.goto(confirmation)
            if "expired" not in (await page.locator("body").inner_text()).lower():
                pytest.fail(
                    "Issuer did not reject the expired email action", pytrace=False
                )
            async with engine_admin.connect() as conn:
                assert (
                    await conn.scalar(
                        text("SELECT count(*) FROM usuarios WHERE email=:email"),
                        {"email": email},
                    )
                    == 0
                )
            await page.goto(backend.FRONT + "/confirmar-email")
            await expect(
                page.get_by_role("heading", name="Confirmar e-mail", exact=True)
            ).to_be_visible()
            await expect(
                page.get_by_role("button", name="Continuar confirmação")
            ).to_be_enabled()
            assert (await browser_request(page, "/auth/session"))["status"] == 401
            await page.get_by_text("Não conseguiu continuar?").click()
            await expect(
                page.get_by_text("Se o link expirou", exact=False)
            ).to_be_visible()
            await page.get_by_role("link", name="Recomeçar cadastro").click()
            await expect(
                page.get_by_role("heading", name="Criar conta", exact=True)
            ).to_be_visible()
            await browser.close()
    except BrowserError:
        pytest.fail("Expired account action recovery did not complete", pytrace=False)
